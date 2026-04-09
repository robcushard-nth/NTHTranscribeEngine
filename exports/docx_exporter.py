"""
NTH Transcript Engine — Word Document (DOCX) Exporter
Sprint 4 | Owner: Backend Agent

Pure function — constructs the document in memory and returns bytes.
No file I/O, no API, no repository access.
"""

import io

from models.transcript import Transcript


def export_docx(transcript: Transcript) -> bytes:
    """Render a Transcript as a .docx file and return raw bytes."""
    from docx import Document  # imported lazily so unrelated tests don't require docx

    doc = Document()
    doc.add_heading("Transcript", level=1)

    subtitle_parts = []
    if transcript.language_detected:
        subtitle_parts.append(f"Language: {transcript.language_detected}")
    if transcript.whisper_model_used:
        subtitle_parts.append(f"Model: {transcript.whisper_model_used}")
    subtitle_parts.append(f"Words: {transcript.word_count}")
    subtitle_parts.append(f"Segments: {transcript.segment_count}")
    doc.add_paragraph(" · ".join(subtitle_parts))

    for seg in transcript.segments:
        para = doc.add_paragraph()
        ts_run = para.add_run(f"[{seg.formatted_start}] ")
        ts_run.bold = True
        para.add_run(seg.text.strip())

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
